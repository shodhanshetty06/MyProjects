"""Unit tests for `app.modules.ats_resume_generator.docx_renderer` — DOCX
generation, no DB/HTTP/AI I-O required.
"""
from __future__ import annotations

import io

from docx import Document

from app.modules.ats_resume_generator.content import ResumeContent
from app.modules.ats_resume_generator.docx_renderer import render_docx
from app.modules.ats_resume_generator.templates import TEMPLATES

_CONTENT = ResumeContent(
    contact={
        "full_name": "Jane Doe",
        "phone": "+1 555-0100",
        "email": "jane@example.com",
        "linkedin_url": "linkedin.com/in/janedoe",
        "github_url": "github.com/janedoe",
        "portfolio_url": "janedoe.dev",
        "location": "Remote",
    },
    summary="Backend engineer with 5 years of experience building APIs.",
    grouped_skills=[
        {"category": "Languages", "skills": ["Python", "SQL"]},
        {"category": "Tools", "skills": ["Docker"]},
    ],
    experience=[
        {
            "title": "Senior Backend Engineer",
            "company": "Acme Corp",
            "location": "Remote",
            "start_date": "2021",
            "end_date": None,
            "is_current": True,
            "description": "Own the payments platform.",
            "highlights": ["Cut checkout latency by 30%", "Led a team of 4 engineers"],
        }
    ],
    projects=[
        {
            "name": "Order Service",
            "description": "Rebuilt the order pipeline.",
            "technologies": ["Kafka", "Python"],
            "highlights": ["Reduced order failures by 20%", "Cut p99 latency by 40%", "Onboarded 3 engineers"],
        }
    ],
    education=[
        {
            "institution": "State University",
            "degree": "B.Sc.",
            "field_of_study": "Computer Science",
            "start_date": "2015",
            "end_date": "2019",
        }
    ],
    certifications=[{"name": "AWS Certified Developer", "issuer": "AWS"}],
    achievements=["Speaker at PyCon 2023"],
    languages=[{"language": "English", "proficiency": "native"}],
)


def _extract_text(docx_bytes: bytes) -> str:
    document = Document(io.BytesIO(docx_bytes))
    return "\n".join(p.text for p in document.paragraphs)


def test_render_docx_produces_a_valid_zip_document() -> None:
    data = render_docx(_CONTENT, TEMPLATES["modern_ats"])
    assert data[:2] == b"PK"  # DOCX is a zip archive
    Document(io.BytesIO(data))  # round-trips without raising


def test_render_docx_includes_header_and_all_sections_in_order() -> None:
    # `modern_ats` uses the "uppercase-rule" heading style, so section
    # headings render fully upper-cased (e.g. "LANGUAGES") — distinct from
    # mixed-case body text like the "Languages: Python, SQL" skill-category
    # label, so matching on the exact upper-cased heading avoids false
    # substring collisions with body content.
    text = _extract_text(render_docx(_CONTENT, TEMPLATES["modern_ats"]))

    assert "Jane Doe" in text
    assert "jane@example.com" in text
    assert "PROFESSIONAL SUMMARY" in text
    assert _CONTENT.summary in text
    assert "TECHNICAL SKILLS" in text
    assert "Python" in text
    assert "PROFESSIONAL EXPERIENCE" in text
    assert "Senior Backend Engineer" in text
    assert "Cut checkout latency by 30%" in text
    assert "PROJECTS" in text
    assert "Order Service" in text
    assert "EDUCATION" in text
    assert "State University" in text
    assert "CERTIFICATIONS" in text
    assert "AWS Certified Developer" in text
    assert "ACHIEVEMENTS" in text
    assert "Speaker at PyCon 2023" in text
    assert "LANGUAGES" in text
    assert "English" in text

    order = [
        text.index("PROFESSIONAL SUMMARY"),
        text.index("TECHNICAL SKILLS"),
        text.index("PROFESSIONAL EXPERIENCE"),
        text.index("PROJECTS"),
        text.index("EDUCATION"),
        text.index("CERTIFICATIONS"),
        text.index("ACHIEVEMENTS"),
        text.index("LANGUAGES"),
    ]
    assert order == sorted(order)


def test_render_docx_omits_empty_sections() -> None:
    minimal = ResumeContent(contact={"full_name": "Solo Person"}, summary="")
    text = _extract_text(render_docx(minimal, TEMPLATES["minimal"]))

    assert "Professional Summary" not in text
    assert "Technical Skills" not in text
    assert "Professional Experience" not in text


def test_render_docx_applies_template_margins() -> None:
    from docx.shared import Inches

    data = render_docx(_CONTENT, TEMPLATES["minimal"])
    document = Document(io.BytesIO(data))
    section = document.sections[0]

    assert section.left_margin == Inches(TEMPLATES["minimal"].margins_inches)


def test_render_docx_works_for_every_template() -> None:
    for template in TEMPLATES.values():
        data = render_docx(_CONTENT, template)
        assert data[:2] == b"PK"
