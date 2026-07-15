"""Unit tests for `app.modules.ats_resume_generator.documents` — DOCX/PDF
generation, no DB/HTTP/AI I-O required.
"""
from __future__ import annotations

import io

from docx import Document

from app.modules.ats_resume_generator.documents import ATSDocumentContent, generate_docx, generate_pdf

_CONTENT = ATSDocumentContent(
    title="Jane Doe — Backend Engineer Resume",
    summary="Backend engineer with 5 years of experience building APIs.",
    skills=["Python", "PostgreSQL", "Docker"],
    experience=[
        {
            "title": "Senior Backend Engineer",
            "company": "Acme Corp",
            "location": "Remote",
            "start_date": "2021",
            "end_date": None,
            "is_current": True,
            "description": "Own the payments platform.",
            "highlights": ["Cut checkout latency by 30%", "Led a team of 4 engineers"],
        }
    ],
    projects=[
        {
            "name": "Order Service",
            "description": "Rebuilt the order pipeline.",
            "technologies": ["Kafka", "Python"],
            "highlights": ["Reduced order failures by 20%"],
        }
    ],
    education=[
        {
            "institution": "State University",
            "degree": "B.Sc.",
            "field_of_study": "Computer Science",
            "start_date": "2015",
            "end_date": "2019",
        }
    ],
    certifications=[{"name": "AWS Certified Developer", "issuer": "AWS"}],
    languages=[{"language": "English", "proficiency": "native"}],
    achievements=["Speaker at PyCon 2023"],
)


def test_generate_docx_produces_valid_document_with_expected_sections() -> None:
    data = generate_docx(_CONTENT)
    assert data[:2] == b"PK"  # DOCX is a zip archive

    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert _CONTENT.title in text
    assert "Summary" in text
    assert _CONTENT.summary in text
    assert "Python" in text
    assert "Senior Backend Engineer" in text
    assert "Cut checkout latency by 30%" in text
    assert "Order Service" in text
    assert "AWS Certified Developer" in text
    assert "Speaker at PyCon 2023" in text


def test_generate_docx_omits_empty_sections() -> None:
    minimal = ATSDocumentContent(title="Minimal Resume", summary="")
    data = generate_docx(minimal)
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for p in document.paragraphs)

    assert "Summary" not in text
    assert "Skills" not in text
    assert "Experience" not in text


def test_generate_pdf_produces_pdf_bytes() -> None:
    data = generate_pdf(_CONTENT)
    assert data[:5] == b"%PDF-"
    assert len(data) > 100


def test_generate_pdf_escapes_special_characters() -> None:
    tricky = ATSDocumentContent(
        title="R&D <Engineer>",
        summary="Worked on A&B < C > D & used <script>alert(1)</script> safely.",
    )
    data = generate_pdf(tricky)
    assert data[:5] == b"%PDF-"
