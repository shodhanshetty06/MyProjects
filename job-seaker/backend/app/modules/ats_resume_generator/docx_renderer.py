"""ATS Resume Generator module — DOCX rendering.

Single-column, no tables/text-boxes/images layout (ATS parsers handle this
most reliably), parameterized by a `TemplateSpec` (`templates.py`) so the six
premium templates share one renderer and differ only in typography/color/
spacing data. The PDF is generated *from* this DOCX (see `pdf_converter.py`)
rather than as an independent pass, so the two formats can never drift.
"""
from __future__ import annotations

from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.modules.ats_resume_generator.content import ResumeContent
from app.modules.ats_resume_generator.templates import TemplateSpec


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _join(parts: list[str | None], sep: str) -> str:
    return sep.join(p for p in parts if p)


def _set_bottom_border(paragraph, color: str, size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color.lstrip("#").upper())
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _experience_header(item: dict) -> tuple[str, str]:
    left = _join([item.get("title"), item.get("company")], " — ")
    if item.get("location"):
        left = _join([left, item["location"]], " · ")
    end = item.get("end_date") or ("Present" if item.get("is_current") else None)
    right = _join([item.get("start_date"), end], " – ")
    return left, right


def _education_line(item: dict) -> str:
    header = _join([item.get("degree"), item.get("field_of_study")], " in ")
    if item.get("institution"):
        header = f"{header}, {item['institution']}" if header else item["institution"]
    date_range = _join([item.get("start_date"), item.get("end_date")], " – ")
    if item.get("gpa"):
        header = f"{header} (GPA: {item['gpa']})" if header else f"GPA: {item['gpa']}"
    return f"{header} ({date_range})" if date_range and header else (header or date_range)


def _certification_line(item: dict) -> str:
    parts = [_join([item.get("name"), item.get("issuer")], " — ")]
    if item.get("issue_date"):
        parts.append(f"({item['issue_date']})")
    return " ".join(p for p in parts if p)


def _language_line(item: dict) -> str:
    if item.get("proficiency"):
        return f"{item.get('language', '')} ({item['proficiency']})"
    return item.get("language", "")


class _DocxBuilder:
    def __init__(self, document: Document, template: TemplateSpec) -> None:
        self.document = document
        self.template = template

    def heading(self, text: str) -> None:
        style = self.template.heading_style
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(self.template.section_spacing_pt)
        paragraph.paragraph_format.space_after = Pt(4)
        display_text = text.upper() if style == "uppercase-rule" else text
        run = paragraph.add_run(display_text)
        run.font.name = self.template.heading_font
        run.font.size = Pt(self.template.base_font_size_pt + 1.5)
        run.font.bold = True
        run.font.color.rgb = _rgb(self.template.heading_color)
        if style == "small-caps":
            run.font.small_caps = True
        if style in ("uppercase-rule", "bold-line") and self.template.accent_color:
            _set_bottom_border(paragraph, self.template.accent_color)

    def body(self, text: str) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = self.template.line_spacing
        run = paragraph.add_run(text)
        run.font.name = self.template.body_font
        run.font.size = Pt(self.template.base_font_size_pt)

    def bullet(self, text: str) -> None:
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = self.template.line_spacing
        run = paragraph.add_run(f"{self.template.bullet_char}  {text}")
        run.font.name = self.template.body_font
        run.font.size = Pt(self.template.base_font_size_pt)

    def two_column_line(self, left: str, right: str, *, bold_left: bool = True) -> None:
        """Left-aligned label + right-aligned date range via a tab stop —
        never a table, so ATS text extraction stays linear."""
        paragraph = self.document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(1)
        section_width = self.document.sections[0].page_width - self.document.sections[0].left_margin - self.document.sections[0].right_margin
        paragraph.paragraph_format.tab_stops.add_tab_stop(section_width, WD_TAB_ALIGNMENT.RIGHT)
        left_run = paragraph.add_run(left)
        left_run.font.name = self.template.body_font
        left_run.font.size = Pt(self.template.base_font_size_pt)
        left_run.font.bold = bold_left
        if right:
            right_run = paragraph.add_run(f"\t{right}")
            right_run.font.name = self.template.body_font
            right_run.font.size = Pt(self.template.base_font_size_pt)
            right_run.italic = True


def render_docx(content: ResumeContent, template: TemplateSpec) -> bytes:
    document = Document()
    section = document.sections[0]
    margin = Inches(template.margins_inches)
    section.left_margin = section.right_margin = margin
    section.top_margin = section.bottom_margin = margin

    normal = document.styles["Normal"]
    normal.font.name = template.body_font
    normal.font.size = Pt(template.base_font_size_pt)
    normal.paragraph_format.line_spacing = template.line_spacing

    builder = _DocxBuilder(document, template)

    # --- Header: name + contact line ---------------------------------------
    contact = content.contact
    name_paragraph = document.add_paragraph()
    name_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_paragraph.paragraph_format.space_after = Pt(2)
    name_run = name_paragraph.add_run(contact.get("full_name") or "")
    name_run.font.name = template.heading_font
    name_run.font.size = Pt(template.base_font_size_pt + 9)
    name_run.font.bold = True
    name_run.font.color.rgb = _rgb(template.heading_color)

    contact_line = _join(
        [
            contact.get("phone"),
            contact.get("email"),
            contact.get("linkedin_url"),
            contact.get("github_url"),
            contact.get("portfolio_url"),
            contact.get("location"),
        ],
        "  |  ",
    )
    if contact_line:
        contact_paragraph = document.add_paragraph()
        contact_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        contact_paragraph.paragraph_format.space_after = Pt(6)
        contact_run = contact_paragraph.add_run(contact_line)
        contact_run.font.name = template.body_font
        contact_run.font.size = Pt(template.base_font_size_pt - 1)
        if template.accent_color:
            _set_bottom_border(contact_paragraph, template.accent_color)

    # --- Professional Summary -----------------------------------------------
    if content.summary:
        builder.heading("Professional Summary")
        builder.body(content.summary)

    # --- Technical Skills (grouped) ------------------------------------------
    if content.grouped_skills:
        builder.heading("Technical Skills")
        for group in content.grouped_skills:
            skills = group.get("skills") or []
            if not skills:
                continue
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2)
            label_run = paragraph.add_run(f"{group.get('category', 'Skills')}: ")
            label_run.font.bold = True
            label_run.font.name = template.body_font
            label_run.font.size = Pt(template.base_font_size_pt)
            value_run = paragraph.add_run(", ".join(skills))
            value_run.font.name = template.body_font
            value_run.font.size = Pt(template.base_font_size_pt)

    # --- Professional Experience (reverse chronological, as provided) -------
    if content.experience:
        builder.heading("Professional Experience")
        for item in content.experience:
            left, right = _experience_header(item)
            builder.two_column_line(left, right)
            if item.get("description"):
                builder.body(item["description"])
            for highlight in item.get("highlights") or []:
                builder.bullet(highlight)

    # --- Projects -------------------------------------------------------------
    if content.projects:
        builder.heading("Projects")
        for item in content.projects:
            name = item.get("name", "")
            technologies = ", ".join(item.get("technologies") or [])
            builder.two_column_line(name, "", bold_left=True)
            if technologies:
                builder.body(f"Technologies: {technologies}")
            if item.get("description"):
                builder.body(item["description"])
            for highlight in item.get("highlights") or []:
                builder.bullet(highlight)

    # --- Education --------------------------------------------------------------
    if content.education:
        builder.heading("Education")
        for item in content.education:
            builder.bullet(_education_line(item))

    # --- Certifications -----------------------------------------------------
    if content.certifications:
        builder.heading("Certifications")
        for item in content.certifications:
            builder.bullet(_certification_line(item))

    # --- Achievements -------------------------------------------------------
    if content.achievements:
        builder.heading("Achievements")
        for item in content.achievements:
            builder.bullet(item)

    # --- Languages ------------------------------------------------------------
    if content.languages:
        builder.heading("Languages")
        builder.body(", ".join(_language_line(item) for item in content.languages))

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


__all__ = ["render_docx"]
