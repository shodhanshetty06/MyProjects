"""DOCX / PDF generation for optimized ATS resumes.

Deliberately plain, single-column formatting (headings + paragraphs + bullet
lists — no tables, text boxes, columns, or images) since that's the layout
ATS parsers handle most reliably; this module intentionally does not try to
visually match a "designed" resume template.

Pure functions over an `ATSDocumentContent` dataclass so both formats are
built from the same section data and are trivial to unit test without S3 or
a DB.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


@dataclass
class ATSDocumentContent:
    title: str
    summary: str
    skills: list[str] = field(default_factory=list)
    experience: list[dict] = field(default_factory=list)
    projects: list[dict] = field(default_factory=list)
    education: list[dict] = field(default_factory=list)
    certifications: list[dict] = field(default_factory=list)
    languages: list[dict] = field(default_factory=list)
    achievements: list[str] = field(default_factory=list)


def _join(parts: list[str | None], sep: str) -> str:
    return sep.join(p for p in parts if p)


def _experience_line(item: dict) -> str:
    header = _join([item.get("title"), item.get("company")], " — ")
    if item.get("location"):
        header = _join([header, item["location"]], " · ")
    end = item.get("end_date") or ("Present" if item.get("is_current") else None)
    date_range = _join([item.get("start_date"), end], " – ")
    if date_range:
        header = f"{header} ({date_range})" if header else date_range
    return header


def _education_line(item: dict) -> str:
    header = _join([item.get("degree"), item.get("field_of_study")], " in ")
    if item.get("institution"):
        header = f"{header}, {item['institution']}" if header else item["institution"]
    date_range = _join([item.get("start_date"), item.get("end_date")], " – ")
    if date_range:
        header = f"{header} ({date_range})" if header else date_range
    return header


def _certification_line(item: dict) -> str:
    return _join([item.get("name"), item.get("issuer")], " — ")


def _language_line(item: dict) -> str:
    if item.get("proficiency"):
        return f"{item.get('language', '')} ({item['proficiency']})"
    return item.get("language", "")


def generate_docx(content: ATSDocumentContent) -> bytes:
    document = Document()
    document.styles["Normal"].font.name = "Calibri"
    document.styles["Normal"].font.size = Pt(11)

    document.add_heading(content.title, level=0)

    def add_section(heading: str) -> None:
        document.add_heading(heading, level=1)

    if content.summary:
        add_section("Summary")
        document.add_paragraph(content.summary)

    if content.skills:
        add_section("Skills")
        document.add_paragraph(", ".join(content.skills))

    if content.experience:
        add_section("Experience")
        for item in content.experience:
            document.add_paragraph(_experience_line(item), style="Heading 2")
            if item.get("description"):
                document.add_paragraph(item["description"])
            for highlight in item.get("highlights") or []:
                document.add_paragraph(highlight, style="List Bullet")

    if content.projects:
        add_section("Projects")
        for item in content.projects:
            header = item.get("name", "")
            if item.get("technologies"):
                header = f"{header} ({', '.join(item['technologies'])})"
            document.add_paragraph(header, style="Heading 2")
            if item.get("description"):
                document.add_paragraph(item["description"])
            for highlight in item.get("highlights") or []:
                document.add_paragraph(highlight, style="List Bullet")

    if content.education:
        add_section("Education")
        for item in content.education:
            document.add_paragraph(_education_line(item), style="List Bullet")

    if content.certifications:
        add_section("Certifications")
        for item in content.certifications:
            document.add_paragraph(_certification_line(item), style="List Bullet")

    if content.languages:
        add_section("Languages")
        document.add_paragraph(", ".join(_language_line(item) for item in content.languages))

    if content.achievements:
        add_section("Achievements")
        for item in content.achievements:
            document.add_paragraph(item, style="List Bullet")

    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def generate_pdf(content: ATSDocumentContent) -> bytes:
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        title=content.title,
    )
    styles = getSampleStyleSheet()
    heading_style = ParagraphStyle(
        "ATSHeading", parent=styles["Heading2"], spaceBefore=12, spaceAfter=4
    )
    subheading_style = ParagraphStyle(
        "ATSSubheading", parent=styles["Heading3"], spaceBefore=8, spaceAfter=2
    )
    body_style = ParagraphStyle("ATSBody", parent=styles["Normal"], spaceAfter=4)
    bullet_style = ParagraphStyle(
        "ATSBullet", parent=styles["Normal"], leftIndent=14, spaceAfter=2
    )

    def p(text: str, style: ParagraphStyle) -> Paragraph:
        return Paragraph(escape(text), style)

    story: list = [p(content.title, styles["Title"]), Spacer(1, 8)]

    def add_section(heading: str) -> None:
        story.append(p(heading, heading_style))

    if content.summary:
        add_section("Summary")
        story.append(p(content.summary, body_style))

    if content.skills:
        add_section("Skills")
        story.append(p(", ".join(content.skills), body_style))

    if content.experience:
        add_section("Experience")
        for item in content.experience:
            story.append(p(_experience_line(item), subheading_style))
            if item.get("description"):
                story.append(p(item["description"], body_style))
            for highlight in item.get("highlights") or []:
                story.append(p(f"• {highlight}", bullet_style))

    if content.projects:
        add_section("Projects")
        for item in content.projects:
            header = item.get("name", "")
            if item.get("technologies"):
                header = f"{header} ({', '.join(item['technologies'])})"
            story.append(p(header, subheading_style))
            if item.get("description"):
                story.append(p(item["description"], body_style))
            for highlight in item.get("highlights") or []:
                story.append(p(f"• {highlight}", bullet_style))

    if content.education:
        add_section("Education")
        for item in content.education:
            story.append(p(f"• {_education_line(item)}", bullet_style))

    if content.certifications:
        add_section("Certifications")
        for item in content.certifications:
            story.append(p(f"• {_certification_line(item)}", bullet_style))

    if content.languages:
        add_section("Languages")
        story.append(p(", ".join(_language_line(item) for item in content.languages), body_style))

    if content.achievements:
        add_section("Achievements")
        for item in content.achievements:
            story.append(p(f"• {item}", bullet_style))

    doc.build(story)
    return buffer.getvalue()


__all__ = ["ATSDocumentContent", "generate_docx", "generate_pdf"]
